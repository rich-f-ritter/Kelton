#!/usr/bin/env python3
"""Build the branded 'Rent Comps - High Level' Excel workbook from a deal config
+ a HelloData unit-details CSV.

Usage:
    python build_workbook.py config.json [out.xlsx]

Sourcing (see ../reference/methodology.md):
  - Per-bedroom RENT  = HelloData T{t_days}-day EXECUTED asking (market) rent
                        (leases that went off-market in the trailing window).
  - Total-Effective   = unit-weighted HelloData effective rent.
  - SF (per-bed+total)= CoStar avg SF by unit type / property avg (from config).
  - Unit counts       = CoStar (comps) / rent roll (subject), from config.
  - Occ %             = CoStar occupied (config);  Leased % = HelloData (computed).
  - Rent & PSF columns get a color-scale heat (low=pale red -> high=pale blue),
    scaled down each column across the comps; comps grouped by submarket cluster.
"""
import sys, json
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import ColorScaleRule

NAVY="074070"; HDR="DCE6F1"; SUBJ="FFF6E6"; GRID="D9D9D9"; WHITE="FFFFFF"
CLUSTER_TINTS=["EAF1F8","EEEFF2","ECF3EE","F3EEF5","FBF4E6","EAF3F3"]  # cool, subtle, cycles

def hd_metrics(df, key, t90, asof):
    s = df[df["Property Name"] == key]
    exe = s[(s["off"] >= t90) & (s["off"] <= asof)]
    out = {}
    for b in [0, 1, 2, 3]:
        e = exe[exe["Bedrooms"] == b]
        if len(e):
            out[b] = dict(ask=round(e["Last Asking Rent"].mean()),
                          eff=round(e["Last Effective Rent"].mean()))
    idx = s.sort_values("on").groupby("Unit").tail(1); tot = s["Unit"].nunique()
    out["leased"] = (1 - idx["off"].isna().sum()/tot) if tot else None
    return out

def blended(counts, hd, field):
    n = d = 0.0
    for i, b in enumerate([0,1,2,3]):
        if b in hd and counts[i] > 0:
            n += hd[b][field]*counts[i]; d += counts[i]
    return round(n/d) if d else None

def main(cfg_path, out_path=None):
    cfg = json.load(open(cfg_path))
    out_path = out_path or "Rent_Comps_High_Level.xlsx"
    df = pd.read_csv(cfg["hellodata_csv"])
    df["off"] = pd.to_datetime(df["Off Market Date"], errors="coerce")
    df["on"]  = pd.to_datetime(df["On Market Date"], errors="coerce")
    asof = pd.Timestamp(cfg.get("as_of", "today"))
    t90 = asof - pd.Timedelta(days=cfg.get("t_days", 90))

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Rent Comps - High Level"
    ws.sheet_view.showGridLines = False
    thin = Side(style="thin", color=GRID); border = Border(thin, thin, thin, thin)
    SECT = Side(style="medium", color="7F7F7F"); CLDIV = Side(style="thin", color="9AA7B4")

    def C(r, c, v=None, *, bold=False, color="404040", size=10, fill=None,
          align="center", fmt=None, wrap=False, italic=False, serif=False):
        cc = ws.cell(r, c, v)
        cc.font = Font(name="Cambria" if serif else "Calibri", bold=bold, color=color, size=size, italic=italic)
        cc.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        if fill: cc.fill = PatternFill("solid", fgColor=fill)
        if fmt: cc.number_format = fmt
        cc.border = border; return cc

    title = cfg.get("title", "Rent Comps")
    ws.merge_cells("A1:AB1")
    C(1,1,f"{title.upper()}  —  RENT COMPS (HIGH LEVEL)   ·   {len(cfg['comps'])} comps, HelloData T{cfg.get('t_days',90)}-day executed rents",
      bold=True, color=WHITE, size=13, fill=NAVY, align="left", serif=True); ws.row_dimensions[1].height=26
    ws.merge_cells("A2:AB2")
    C(2,1,f"Rent (by bedroom) = HelloData T{cfg.get('t_days',90)}-day EXECUTED asking (market) rent. SF (by bedroom & total) = CoStar avg SF by unit type "
          "(subject SF/counts = rent roll). Comp unit counts = CoStar. Occ % = CoStar occupied; Leased % = HelloData; "
          "Conc = 1 − Eff ÷ Gross; PSF = total rent ÷ total SF.",
      italic=True, color="6F6F70", size=9, align="left", wrap=True); ws.row_dimensions[2].height=38

    GH, H, C0 = 4, 5, 11
    idcols=[("#",4),("Property",27),("Group / Submarket",17),("Owner",17),("Yr\nBuilt",6),
            ("Dist\n(mi)",6),("Occ %\n(CoStar)",8),("Leased %\n(HD)",8),("Conc.\n(impl.)",7),("Units",6)]
    for j,(lbl,w) in enumerate(idcols,1):
        ws.column_dimensions[L(j)].width=w
        C(GH,j,"",fill=HDR); C(H,j,lbl,bold=True,color=NAVY,size=9,fill=HDR,wrap=True)
    blocks=["Studio","1 BR","2 BR","3 BR","Total — Gross","Total — Effective"]
    subc={"Studio":["#","SF","Rent"],"1 BR":["#","SF","Rent"],"2 BR":["#","SF","Rent"],"3 BR":["#","SF","Rent"],
          "Total — Gross":["#","SF","Rent","PSF"],"Total — Effective":["Rent","PSF"]}
    col=C0; bc={}
    for b in blocks:
        sc=subc[b]; start=col
        ws.merge_cells(start_row=GH,start_column=start,end_row=GH,end_column=start+len(sc)-1)
        C(GH,start,b,bold=True,color=NAVY,size=9.5,fill=HDR)
        for k,s in enumerate(sc): C(H,start+k,s,bold=True,color=NAVY,size=9,fill=HDR)
        bc[b]=(start,sc); col+=len(sc)
    LAST=col-1
    def cof(b,s): st,sc=bc[b]; return st+sc.index(s)
    for c in range(C0,LAST+1): ws.column_dimensions[L(c)].width=6.5
    ws.row_dimensions[GH].height=16; ws.row_dimensions[H].height=26

    def row(r, num, rec, subject=False):
        counts = rec["counts"]; bedsf = rec.get("bed_sf",[None]*4); avgsf = rec["avg_sf"]
        hd = hd_metrics(df, rec["hd_key"], t90, asof)
        bg = SUBJ if subject else None
        C(r,1,num if num!="" else "",bold=True,color=NAVY,fill=bg)
        C(r,2,rec["name"],bold=subject,color="1A2B3C",align="left",fill=bg)
        C(r,3,rec.get("cluster","SUBJECT") if subject else rec.get("cluster",""),color="404040",size=9,fill=bg)
        C(r,4,rec.get("owner",""),color="404040",size=9,align="left",fill=bg)
        C(r,5,rec.get("year",""),fill=bg)
        C(r,6,rec.get("dist",""),fill=bg,fmt="0.0")
        C(r,7,rec.get("occ_costar"),fill=bg,fmt="0.0%")
        C(r,8,hd.get("leased"),fill=bg,fmt="0.0%")
        for i,(bl,bi) in enumerate([("Studio",0),("1 BR",1),("2 BR",2),("3 BR",3)]):
            cc=cof(bl,"#"); cs=cof(bl,"SF"); cr=cof(bl,"Rent")
            C(r,cc,counts[i] if counts[i] else "",fill=bg)
            C(r,cs,(bedsf[i] if (counts[i] and bedsf[i]) else None),fill=bg,fmt="#,##0")
            ask=hd[bi]["ask"] if (bi in hd and counts[i]>0) else None
            C(r,cr,ask,fill=bg,fmt="$#,##0")
        units=rec.get("units", sum(counts)); g_ask=blended(counts,hd,"ask"); g_eff=blended(counts,hd,"eff")
        tgc=cof("Total — Gross","#"); tgs=cof("Total — Gross","SF"); tgr=cof("Total — Gross","Rent"); tgp=cof("Total — Gross","PSF")
        C(r,tgc,units,bold=True,color=NAVY,fill=bg)
        C(r,tgs,avgsf,bold=True,color=NAVY,fill=bg,fmt="#,##0")
        C(r,tgr,g_ask,bold=True,color=NAVY,fill=bg,fmt="$#,##0")
        C(r,tgp,f"=IF(N({L(tgs)}{r})>0,{L(tgr)}{r}/{L(tgs)}{r},\"\")",bold=True,color=NAVY,fill=bg,fmt="$#,##0.00")
        ter=cof("Total — Effective","Rent"); tep=cof("Total — Effective","PSF")
        C(r,ter,g_eff,bold=True,color=NAVY,fill=bg,fmt="$#,##0")
        C(r,tep,f"=IF(N({L(tgs)}{r})>0,{L(ter)}{r}/{L(tgs)}{r},\"\")",bold=True,color=NAVY,fill=bg,fmt="$#,##0.00")
        C(r,9,((1-g_eff/g_ask) if (g_ask and g_eff) else None),fill=bg,fmt="0%")
        C(r,10,units,fill=bg)

    row(6,"S",cfg["subject"],subject=True)
    # Optional: overwrite subject-row cells with live model-link formulas so the tab
    # drops into the TMG model with row 6 tied to its Assumptions / Cash Flow tabs
    # (values resolve inside the model; standalone they show #REF!). Keys are cell
    # addresses (e.g. "N6"); a value may be a formula string or {"array": "=..."}.
    scf = cfg.get("subject_cell_formulas")
    if scf:
        from openpyxl.worksheet.formula import ArrayFormula
        for addr, fml in scf.items():
            cell = ws[addr]
            cell.value = ArrayFormula(addr, fml["array"]) if isinstance(fml, dict) else fml
    C(8,1,"Rent Comps",bold=True,color=NAVY,size=11,align="left")
    for j in range(2,LAST+1): C(8,j,"")
    r0=9
    for i,rec in enumerate(cfg["comps"]): row(r0+i, i+1, rec)
    TOT=r0+len(cfg["comps"])

    # comp total / weighted-average row
    C(TOT,1,"",fill=HDR); C(TOT,2,"Comp Total / Wtd Avg",bold=True,color=NAVY,align="left",fill=HDR)
    for j in range(3,11): C(TOT,j,"",fill=HDR)
    C(TOT,10,f"=SUM(J9:J{TOT-1})",bold=True,color=NAVY,fill=HDR)
    for bl in ["Studio","1 BR","2 BR","3 BR"]:
        cc=cof(bl,"#"); cs=cof(bl,"SF"); cr=cof(bl,"Rent")
        C(TOT,cc,f"=SUM({L(cc)}9:{L(cc)}{TOT-1})",bold=True,color=NAVY,fill=HDR)
        C(TOT,cs,f"=IFERROR(SUMPRODUCT({L(cs)}9:{L(cs)}{TOT-1},{L(cc)}9:{L(cc)}{TOT-1})/SUMPRODUCT(({L(cs)}9:{L(cs)}{TOT-1}>0)*{L(cc)}9:{L(cc)}{TOT-1}),\"\")",bold=True,color=NAVY,fill=HDR,fmt="#,##0")
        C(TOT,cr,f"=IFERROR(SUMPRODUCT({L(cr)}9:{L(cr)}{TOT-1},{L(cc)}9:{L(cc)}{TOT-1})/SUMPRODUCT(({L(cr)}9:{L(cr)}{TOT-1}>0)*{L(cc)}9:{L(cc)}{TOT-1}),\"\")",bold=True,color=NAVY,fill=HDR,fmt="$#,##0")
    tgc=cof("Total — Gross","#"); tgs=cof("Total — Gross","SF"); tgr=cof("Total — Gross","Rent"); tgp=cof("Total — Gross","PSF")
    C(TOT,tgc,f"=SUM({L(tgc)}9:{L(tgc)}{TOT-1})",bold=True,color=NAVY,fill=HDR)
    C(TOT,tgs,f"=IFERROR(SUMPRODUCT({L(tgs)}9:{L(tgs)}{TOT-1},{L(tgc)}9:{L(tgc)}{TOT-1})/SUM({L(tgc)}9:{L(tgc)}{TOT-1}),\"\")",bold=True,color=NAVY,fill=HDR,fmt="#,##0")
    C(TOT,tgr,f"=IFERROR(SUMPRODUCT({L(tgr)}9:{L(tgr)}{TOT-1},{L(tgc)}9:{L(tgc)}{TOT-1})/SUM({L(tgc)}9:{L(tgc)}{TOT-1}),\"\")",bold=True,color=NAVY,fill=HDR,fmt="$#,##0")
    C(TOT,tgp,f"=IF(N({L(tgs)}{TOT})>0,{L(tgr)}{TOT}/{L(tgs)}{TOT},\"\")",bold=True,color=NAVY,fill=HDR,fmt="$#,##0.00")
    ter=cof("Total — Effective","Rent"); tep=cof("Total — Effective","PSF")
    C(TOT,ter,f"=IFERROR(SUMPRODUCT({L(ter)}9:{L(ter)}{TOT-1},{L(tgc)}9:{L(tgc)}{TOT-1})/SUM({L(tgc)}9:{L(tgc)}{TOT-1}),\"\")",bold=True,color=NAVY,fill=HDR,fmt="$#,##0")
    C(TOT,tep,f"=IF(N({L(tgs)}{TOT})>0,{L(ter)}{TOT}/{L(tgs)}{TOT},\"\")",bold=True,color=NAVY,fill=HDR,fmt="$#,##0.00")
    C(TOT,9,f"=IFERROR(1-{L(ter)}{TOT}/{L(tgr)}{TOT},\"\")",bold=True,color=NAVY,fill=HDR,fmt="0%")

    # section boxes around the 4 bedroom blocks
    def box(c0,c1,a,z):
        for rr in range(a,z+1):
            for cc in range(c0,c1+1):
                cell=ws.cell(rr,cc); b=cell.border
                cell.border=Border(left=SECT if cc==c0 else b.left, right=SECT if cc==c1 else b.right,
                                   top=SECT if rr==a else b.top, bottom=SECT if rr==z else b.bottom)
    for bl in ["Studio","1 BR","2 BR","3 BR"]:
        st,sc=bc[bl]; box(st,st+len(sc)-1,GH,TOT)

    # group comps by submarket cluster: merged label + faint tint + divider
    seen={}; tints={}
    for rec in cfg["comps"]:
        cl=rec.get("cluster","")
        if cl not in tints: tints[cl]=CLUSTER_TINTS[len(tints)%len(CLUSTER_TINTS)]
    i=0
    while i < len(cfg["comps"]):
        cl=cfg["comps"][i].get("cluster","")
        j=i
        while j+1<len(cfg["comps"]) and cfg["comps"][j+1].get("cluster","")==cl: j+=1
        a,z=r0+i, r0+j
        ws.merge_cells(start_row=a,start_column=3,end_row=z,end_column=3)
        gc=ws.cell(a,3); gc.value=cl
        gc.font=Font(name="Calibri",bold=True,color=NAVY,size=9)
        gc.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        gc.fill=PatternFill("solid",fgColor=tints[cl])
        if z != TOT-1:
            for cc in range(1,LAST+1):
                cell=ws.cell(z,cc); b=cell.border
                cell.border=Border(left=b.left,right=b.right,top=b.top,bottom=CLDIV)
        i=j+1

    # heat on rent columns + total PSF
    def heat(c):
        ws.conditional_formatting.add(f"{L(c)}9:{L(c)}{TOT-1}", ColorScaleRule(
            start_type="min", start_color="F6C9BE", mid_type="percentile", mid_value=50,
            mid_color="FFFFFF", end_type="max", end_color="C7DCEF"))
    for bl in ["Studio","1 BR","2 BR","3 BR"]: heat(cof(bl,"Rent"))
    for cc in [cof("Total — Gross","Rent"),cof("Total — Gross","PSF"),
               cof("Total — Effective","Rent"),cof("Total — Effective","PSF")]: heat(cc)

    fr=TOT+2
    ws.merge_cells(start_row=fr,start_column=1,end_row=fr,end_column=LAST)
    C(fr,1, cfg.get("footnote",
        "Rent & total-PSF columns are color-scaled down each column across the comps (pale red = lower, pale blue = higher; subject = cream, excluded). "
        "Rent = HelloData T-day executed asking; SF = CoStar avg SF by unit type; comps grouped by submarket cluster."),
      italic=True,color="6F6F70",size=8,align="left",wrap=True); ws.row_dimensions[fr].height=46
    ws.freeze_panes="K6"
    wb.save(out_path)
    print(f"wrote {out_path}  ({len(cfg['comps'])} comps, rows 9-{TOT-1}, total row {TOT})")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(1)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
